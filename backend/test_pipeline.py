"""pipeline.py olcumu: dosya -> oturum zinciri.

Varsayilan calisma MODELSIZ. Cozumleme adimi disaridan verilen sahte bir
fonksiyonla degistirilir, boylece zincir ve yollar tek basina olculur.

    .\\.venv\\Scripts\\python.exe test_pipeline.py
    .\\.venv\\Scripts\\python.exe test_pipeline.py --model    (gercek model +
                                                    uploads\\cv.pdf, uctan uca)

Gercek output\\ ve uploads\\ klasorlerine yazilmaz; hepsi gecici klasorde.
"""

import json
import shutil
import sys
import tempfile
from pathlib import Path

import commands
import llm
import render_cv
import session

ROOT = Path(__file__).resolve().parent.parent
UPLOADS = ROOT / "uploads"
OUT = ROOT / "output"

gecen = kalan = 0


def kontrol(ad, kosul, ayrinti=""):
    global gecen, kalan
    try:
        ok = bool(kosul)
    except Exception as e:                                   # noqa: BLE001
        ok, ayrinti = False, "{}: {}".format(type(e).__name__, e)
    if ok:
        gecen += 1
        print("  gecti  {}".format(ad))
    else:
        kalan += 1
        print("  KALDI  {}   {}".format(ad, ayrinti))


SATIRLAR = [
    (72, 100, 20, True, "Ayse Seyda Aksakal"),
    (72, 125, 12, False, "Yazilim Gelistirici"),
    (72, 160, 14, True, "IS DENEYIMI"),
    (72, 185, 11, True, "Kidemli Muhendis"),
    (400, 185, 10, False, "2021 - 2024"),
    (86, 205, 10, False, "Olcum altyapisini kurdu ve gecikmeyi dusurdu."),
    (72, 240, 14, True, "EGITIM"),
    (72, 265, 11, False, "Anadolu Universitesi, Bilgisayar Muhendisligi"),
]


def ornek_pdf(path):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    for x, y, punto, kalin, metin in SATIRLAR:
        page.insert_text((x, y), metin, fontsize=punto,
                         fontname="helv" if not kalin else "hebo")
    doc.save(str(path))
    doc.close()


def ornek_docx(path):
    from docx import Document
    from docx.shared import Mm, Pt
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for _, _, punto, kalin, metin in SATIRLAR:
        p = d.add_paragraph()
        r = p.add_run(metin)
        r.bold = kalin
        r.font.size = Pt(punto)
    d.save(str(path))


def sahte_cozumleyici(layout):
    """Kaynaktan okunan gercek metinle CV kurar. Model yok."""
    sp = [s["text"] for p in layout.get("pages", []) for s in p.get("spans", [])]
    blob = " ".join(sp)
    return {
        "name": "Ayse Seyda Aksakal",
        "title": "Yazilim Gelistirici",
        "contact": {"email": "", "phone": "", "location": "", "links": []},
        "sections": [
            {"heading": "IS DENEYIMI", "items": [
                {"title": "Kidemli Muhendis", "subtitle": "", "date": "2021 - 2024",
                 "bullets": [t for t in sp if "gecikmeyi" in t]},
            ]},
            {"heading": "EGITIM", "items": [
                {"title": "Anadolu Universitesi, Bilgisayar Muhendisligi",
                 "subtitle": "", "date": "", "bullets": []},
            ]},
        ],
        "_kaynak_uzunluk": len(blob),
    }, {"sure": 0.0, "thinking": ""}


def test_ayristirici(pipeline, tmp):
    print("\n[1] Ayristirici secimi")

    pdf, docx = tmp / "ornek.pdf", tmp / "ornek.docx"
    ornek_pdf(pdf)
    ornek_docx(docx)

    kontrol("desteklenen uzantilar", pipeline.desteklenen() == [".docx", ".pdf"],
            pipeline.desteklenen())

    d1 = pipeline.ayristir(pdf, tmp / "p1")
    d2 = pipeline.ayristir(docx, tmp / "p2")

    kontrol("PDF parca uretti", len(d1["pages"][0]["spans"]) >= 8,
            len(d1["pages"][0]["spans"]))
    kontrol("DOCX parca uretti", len(d2["pages"][0]["spans"]) >= 8,
            len(d2["pages"][0]["spans"]))
    kontrol("kok anahtarlar ayni", set(d1) == set(d2),
            set(d1) ^ set(d2))
    kontrol("parca alanlari ayni",
            set(d1["pages"][0]["spans"][0]) == set(d2["pages"][0]["spans"][0]))
    kontrol("sayfa alanlari ayni",
            set(d1["pages"][0]) == set(d2["pages"][0]))
    kontrol("iki tarafta da ayni metin",
            "IS DENEYIMI" in " ".join(s["text"] for s in d1["pages"][0]["spans"])
            and "IS DENEYIMI" in " ".join(s["text"] for s in d2["pages"][0]["spans"]))

    kontrol("sabit taban adi kullanildi", (tmp / "p1" / "cv_layout.json").exists())
    kontrol("PDF sayfa goruntusu yazildi", (tmp / "p1" / "cv_page1.png").exists())

    for kotu, beklenen in ((tmp / "x.txt", ValueError),
                           (tmp / "yok.pdf", FileNotFoundError)):
        kotu.write_text("x", encoding="utf-8") if kotu.suffix == ".txt" else None
        try:
            pipeline.ayristir(kotu, tmp / "p3")
            ok = False
            hata = "hata vermedi"
        except Exception as e:                               # noqa: BLE001
            ok = isinstance(e, beklenen)
            hata = type(e).__name__
        kontrol("{} reddedildi".format(kotu.name), ok, hata)


def test_temizle(pipeline, tmp):
    print("\n[2] temizle: semayi Python zorluyor")

    layout = pipeline.ayristir(tmp / "ornek.pdf", tmp / "p4")

    ham = {
        "name": "  Ayse   Seyda  ",
        "unvan_yanlis_anahtar": "atilmali",
        "contact": {"email": "a@b.com", "links": "tek-bag"},
        "sections": [
            {"heading": "IS DENEYIMI", "items": [
                {"title": "Muhendis", "bullets": "tek madde metni"},
                "duz metin kayit",
                {"title": "", "subtitle": "", "date": "", "bullets": []},
            ]},
            {"heading": "", "items": []},
            "bolum degil",
        ],
    }
    cv = pipeline.temizle(ham, layout)

    kontrol("bosluk sadelesti", cv["name"] == "Ayse Seyda", cv["name"])
    kontrol("fazla anahtar atildi", "unvan_yanlis_anahtar" not in cv)
    kontrol("eksik alanlar dolduruldu",
            cv["title"] == "" and cv["contact"]["phone"] == "")
    kontrol("links liste oldu", isinstance(cv["contact"]["links"], list),
            cv["contact"]["links"])
    kontrol("metin bullets listeye cevrildi",
            cv["sections"][0]["items"][0]["bullets"] == ["tek madde metni"])
    kontrol("duz metin kayit sozluk oldu",
            cv["sections"][0]["items"][1]["title"] == "duz metin kayit")
    kontrol("bos kayit atildi", len(cv["sections"][0]["items"]) == 2,
            len(cv["sections"][0]["items"]))
    kontrol("bos bolum atildi", len(cv["sections"]) == 1, len(cv["sections"]))
    kontrol("sozluk olmayan bolum atildi",
            all(isinstance(s, dict) for s in cv["sections"]))
    kontrol("design layout'tan geldi",
            cv["design"]["page"][0] == layout["pages"][0]["width_pt"],
            cv["design"]["page"])
    kontrol("fontlar tasindi", len(cv["design"]["fonts"]) >= 1)

    uzun = {"sections": [{"heading": "A", "items": [
        {"title": "x" * 500, "bullets": ["y" * 900] + ["z"] * 40}]}]}
    cv2 = pipeline.temizle(uzun)
    it = cv2["sections"][0]["items"][0]
    kontrol("uzun baslik kirpildi", len(it["title"]) <= pipeline.MAX_KISA,
            len(it["title"]))
    kontrol("uzun madde kirpildi", len(it["bullets"][0]) <= pipeline.MAX_UZUN,
            len(it["bullets"][0]))
    kontrol("madde sayisi sinirli", len(it["bullets"]) <= pipeline.MAX_MADDE,
            len(it["bullets"]))
    kontrol("bos girdi cokmedi", pipeline.temizle(None)["sections"] == [])


def test_kapsama(pipeline, tmp):
    print("\n[3] kapsama: uydurma yakalaniyor")

    layout = pipeline.ayristir(tmp / "ornek.pdf", tmp / "p5")
    dogru, _ = sahte_cozumleyici(layout)
    oran, eksik = pipeline.kapsama(pipeline.temizle(dogru, layout), layout)
    kontrol("kaynaktan gelen CV tam", oran >= 0.99, "{} {}".format(oran, eksik))
    kontrol("eksik listesi bos", eksik == [], eksik)

    uydurma = {"name": "Baska Biri Tamamen",
               "sections": [{"heading": "HOBILER", "items": [
                   {"title": "Serbest dalis egitmenligi sertifikasi",
                    "bullets": ["Bu cumle kaynak dosyada gecmiyor."]}]}]}
    oran2, eksik2 = pipeline.kapsama(pipeline.temizle(uydurma, layout), layout)
    kontrol("uydurma orani dusuk", oran2 < 0.2, oran2)
    kontrol("eksik parcalar listelendi", len(eksik2) >= 3, eksik2[:2])

    kontrol("noktalama fark etmiyor",
            pipeline.kapsama({"name": "AYSE, SEYDA. AKSAKAL!"}, layout)[0] == 1.0)
    kontrol("bos CV bir dondurur", pipeline.kapsama({}, layout)[0] == 1.0)


def test_benzersiz(pipeline, tmp):
    print("\n[4] Yukleme dosya adi")

    klasor = tmp / "yuk"
    kontrol("uzanti korundu",
            pipeline.benzersiz("Ozgecmis 2024.pdf", klasor).suffix == ".pdf")
    kontrol("bosluk temizlendi",
            " " not in pipeline.benzersiz("Ozgecmis 2024.pdf", klasor).name,
            pipeline.benzersiz("Ozgecmis 2024.pdf", klasor).name)
    kontrol("yol atildi",
            pipeline.benzersiz("../../gizli/cv.pdf", klasor).parent == klasor)
    kontrol("turkce ad cokmedi",
            pipeline.benzersiz("ozgecmis_son.docx", klasor).suffix == ".docx")

    ilk = pipeline.benzersiz("cv.pdf", klasor)
    ilk.write_text("x", encoding="utf-8")
    ikinci = pipeline.benzersiz("cv.pdf", klasor)
    kontrol("cakisma yok", ikinci != ilk, "{} / {}".format(ilk.name, ikinci.name))


def test_zincir(pipeline, tmp):
    print("\n[5] Uctan uca zincir (model yok)")

    s1 = pipeline.calistir(tmp / "ornek.pdf", cozumleyici=sahte_cozumleyici)
    kontrol("oturum acildi", session.var(s1["id"]), s1["id"])
    kontrol("PDF turu bildirildi", s1["tur"] == "pdf", s1["tur"])
    kontrol("bolum sayisi dondu", s1["bolum"] == 2, s1["bolum"])
    kontrol("kayit sayisi dondu", s1["kayit"] == 2, s1["kayit"])
    kontrol("kapsama olculdu", s1["kapsama"] >= 0.99, s1["kapsama"])

    d = session.yol(s1["id"])
    kontrol("cv_layout.json yazildi", (d / "cv_layout.json").exists())
    kontrol("cv_structured.json yazildi", (d / "cv_structured.json").exists())
    kontrol("cv_generated.html yazildi", (d / "cv_generated.html").exists())
    kontrol("cv_page1.png yazildi", (d / "cv_page1.png").exists())
    kontrol("HTML ad iceriyor",
            "Aksakal" in (d / "cv_generated.html").read_text(encoding="utf-8"))
    kontrol("yuklenen dosya kopyalandi",
            (Path(pipeline.UPLOADS) / s1["dosya"]).exists(), s1["dosya"])
    kontrol("meta adi guncellendi", session.meta(s1["id"])["ad"], "")
    kontrol("oturum aktif oldu", session.aktif() == s1["id"])
    kontrol("commands oturuma bagli",
            commands.STRUCT == d / "cv_structured.json")
    kontrol("commands.load calisiyor",
            commands.load()["name"] == "Ayse Seyda Aksakal")

    s2 = pipeline.calistir(tmp / "ornek.docx", cozumleyici=sahte_cozumleyici)
    kontrol("ikinci dosya yeni oturum", s2["id"] != s1["id"],
            "{} / {}".format(s1["id"], s2["id"]))
    kontrol("DOCX turu bildirildi", s2["tur"] == "docx", s2["tur"])
    kontrol("ilk oturum bozulmadi",
            json.loads((d / "cv_structured.json").read_text(encoding="utf-8"))
            ["name"] == "Ayse Seyda Aksakal")

    s3 = pipeline.calistir(tmp / "ornek.pdf", oid=s1["id"],
                           cozumleyici=sahte_cozumleyici)
    kontrol("verilen oturuma yazildi", s3["id"] == s1["id"])
    kontrol("yeni oturum acilmadi", len(session.liste()) == 2,
            len(session.liste()))
    kontrol("ikinci yuklemede ad cakismadi", s3["dosya"] != s1["dosya"],
            "{} / {}".format(s1["dosya"], s3["dosya"]))


def test_hata(pipeline, tmp):
    print("\n[6] Hata halinde temizlik")

    once = len(session.liste())

    def patlayan(layout):
        raise llm.LLMError("model yok")

    try:
        pipeline.calistir(tmp / "ornek.pdf", cozumleyici=patlayan)
        ok = False
    except llm.LLMError:
        ok = True
    kontrol("model hatasi disari verildi", ok)
    kontrol("yarim oturum silindi", len(session.liste()) == once,
            len(session.liste()))

    def bos(layout):
        return {"name": "X", "sections": []}, {}

    try:
        pipeline.calistir(tmp / "ornek.pdf", cozumleyici=bos)
        ok = False
        mesaj = ""
    except ValueError as e:
        ok, mesaj = True, str(e)
    kontrol("bolumsuz sonuc reddedildi", ok, mesaj)
    kontrol("oturum yine silindi", len(session.liste()) == once,
            len(session.liste()))

    bos_pdf = tmp / "bos.pdf"
    import pymupdf
    doc = pymupdf.open()
    doc.new_page(width=595, height=842)
    doc.save(str(bos_pdf))
    doc.close()
    try:
        pipeline.calistir(bos_pdf, cozumleyici=sahte_cozumleyici)
        ok = False
        mesaj = ""
    except ValueError as e:
        ok, mesaj = True, str(e)
    kontrol("metinsiz dosya reddedildi", ok, mesaj)
    kontrol("oturum sayisi degismedi", len(session.liste()) == once,
            len(session.liste()))


def test_model(pipeline):
    print("\n[7] Gercek model + gercek CV")

    kaynak = sorted(list(UPLOADS.glob("*.pdf")) + list(UPLOADS.glob("*.docx")))
    if not kaynak:
        print("  atlandi: uploads klasorunde PDF/DOCX yok")
        return
    if llm.MODEL not in llm.available():
        print("  atlandi: {} kurulu degil".format(llm.MODEL))
        return

    f = kaynak[0]
    print("  kaynak:", f.name, "- model calisiyor, bekle...")
    try:
        s = pipeline.calistir(f)
    except Exception as e:                                   # noqa: BLE001
        kontrol("model zinciri calisti", False, "{}: {}".format(
            type(e).__name__, str(e)[:150]))
        return

    kontrol("model zinciri calisti", s["ok"])
    kontrol("bolum cikti", s["bolum"] >= 3, s["bolum"])
    kontrol("kayit cikti", s["kayit"] >= 5, s["kayit"])
    kontrol("ad bulundu", len(s["ad"]) >= 3, s["ad"])
    kontrol("kapsama %90 uzeri", s["kapsama"] >= 0.90, s["kapsama"])
    kontrol("HTML uretildi", s["uzunluk"] > 1000, s["uzunluk"])
    pipeline.ozet(s)


def main():
    modelli = "--model" in sys.argv

    print("=" * 60)
    print("[0] Modul")
    try:
        import pipeline
    except ImportError as e:
        print("  KALDI  pipeline bulunamadi   {}".format(e))
        print("\n" + "-" * 60)
        print("Sonuc: 0 gecti, 1 kaldi   (temel olcum)")
        return 1

    eksik = [ad for ad in ("ayristir", "cozumle", "temizle", "kapsama",
                           "benzersiz", "calistir")
             if not callable(getattr(pipeline, ad, None))]
    if eksik:
        print("  KALDI  eksik fonksiyon: {}".format(eksik))
        print("\n" + "-" * 60)
        print("Sonuc: 0 gecti, 1 kaldi   (temel olcum)")
        return 1
    print("  gecti  pipeline arayuzu tam")

    tmp = Path(tempfile.mkdtemp(prefix="cvpipe_"))
    eski = (session.KOK, pipeline.UPLOADS,
            commands.OUT, commands.STRUCT, commands.HIST,
            render_cv.OUT, render_cv.STRUCT, render_cv.LAYOUT,
            render_cv.HTML_OUT)
    try:
        session.KOK = tmp / "oturum"
        pipeline.UPLOADS = tmp / "uploads"

        ornek_pdf(tmp / "ornek.pdf")
        ornek_docx(tmp / "ornek.docx")

        test_ayristirici(pipeline, tmp)
        test_temizle(pipeline, tmp)
        test_kapsama(pipeline, tmp)
        test_benzersiz(pipeline, tmp)
        test_zincir(pipeline, tmp)
        test_hata(pipeline, tmp)
        if modelli:
            test_model(pipeline)
        else:
            print("\n[7] Model testi atlandi (--model ile calistir)")
    finally:
        session.KOK = eski[0]
        pipeline.UPLOADS = eski[1]
        (commands.OUT, commands.STRUCT, commands.HIST,
         render_cv.OUT, render_cv.STRUCT, render_cv.LAYOUT,
         render_cv.HTML_OUT) = eski[2:]
        shutil.rmtree(tmp, ignore_errors=True)

    print("\n[8] Gercek klasorler")
    kontrol("commands.STRUCT geri konuldu",
            commands.STRUCT == OUT / "cv_structured.json", commands.STRUCT)
    kontrol("uploads kirletilmedi",
            not (UPLOADS / "ornek.pdf").exists() or modelli)

    print("\n" + "-" * 60)
    print("Sonuc: {} gecti, {} kaldi".format(gecen + 1, kalan))
    return 1 if kalan else 0


if __name__ == "__main__":
    sys.exit(main())
