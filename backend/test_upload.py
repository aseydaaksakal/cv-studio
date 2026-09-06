"""app.py oturumlu uc noktalar olcumu (Asama 6d-7d).

Model YOK. pipeline.cozumle ve classify.classify sahte fonksiyonlarla
degistirilir, boylece /upload ve oturum uclari tek basina olculur.

Gercek output\\ ve uploads\\ klasorlerine YAZILMAZ; session.KOK,
session.OUT ve pipeline.UPLOADS gecici klasore cevrilir, finally icinde
geri konur. app modulu ancak bu cevirmeden SONRA import edilir.

    .\\.venv\\Scripts\\python.exe test_upload.py
    .\\.venv\\Scripts\\python.exe test_upload.py --model   (gercek model turu)
"""

import json
import shutil
import sys
import tempfile
from pathlib import Path

import commands
import pipeline
import render_cv
import session

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"
UPLOADS = ROOT / "uploads"

gecen = kalan = 0


def kontrol(ad, kosul, ayrinti=""):
    global gecen, kalan
    try:
        ok = bool(kosul)
    except Exception as e:                                   # noqa: BLE001
        ok, ayrinti = False, "{}: {}".format(type(e).__name__, e)
    if ok:
        gecen += 1
        print("  gecti  {}".format(ad))
    else:
        kalan += 1
        print("  KALDI  {}   {}".format(ad, ayrinti))


def grup(fn, *a):
    """Bir olcum grubu patlarsa sayi yine de cikmali."""
    try:
        fn(*a)
    except Exception as e:                                   # noqa: BLE001
        kontrol("{} grubu tamamlandi".format(fn.__name__), False,
                "{}: {}".format(type(e).__name__, str(e)[:150]))


# --- ornek dosyalar ---------------------------------------------------

SATIRLAR = [
    (72, 100, 20, True, "Ayse Seyda Aksakal"),
    (72, 125, 12, False, "Yazilim Gelistirici"),
    (72, 160, 14, True, "IS DENEYIMI"),
    (72, 185, 11, True, "Kidemli Muhendis"),
    (400, 185, 10, False, "2021 - 2024"),
    (86, 205, 10, False, "Olcum altyapisini kurdu ve gecikmeyi dusurdu."),
    (72, 240, 14, True, "EGITIM"),
    (72, 265, 11, False, "Anadolu Universitesi, Bilgisayar Muhendisligi"),
]


def ornek_pdf(path):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    for x, y, punto, kalin, metin in SATIRLAR:
        page.insert_text((x, y), metin, fontsize=punto,
                         fontname="hebo" if kalin else "helv")
    doc.save(str(path))
    doc.close()
    return path


def ornek_docx(path):
    from docx import Document
    from docx.shared import Mm, Pt
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for _, _, punto, kalin, metin in SATIRLAR:
        p = d.add_paragraph()
        r = p.add_run(metin)
        r.bold = kalin
        r.font.size = Pt(punto)
    d.save(str(path))
    return path


# --- sahte model ------------------------------------------------------

def sahte_cozumle(layout, model=None):
    """pipeline.cozumle yerine gecer. Kaynaktaki gercek metni kullanir."""
    sp = [s["text"] for p in layout.get("pages", []) for s in p.get("spans", [])]
    return {
        "name": "Ayse Seyda Aksakal",
        "title": "Yazilim Gelistirici",
        "contact": {"email": "", "phone": "", "location": "", "links": []},
        "sections": [
            {"heading": "IS DENEYIMI", "items": [
                {"title": "Kidemli Muhendis", "subtitle": "",
                 "date": "2021 - 2024",
                 "bullets": [t for t in sp if "gecikmeyi" in t]},
            ]},
            {"heading": "EGITIM", "items": [
                {"title": "Anadolu Universitesi, Bilgisayar Muhendisligi",
                 "subtitle": "", "date": "", "bullets": []},
            ]},
        ],
    }, {"sure": 0.0, "thinking": ""}


def sahte_classify(text, cv):
    """classify.classify yerine gecer. Deterministik tek adim."""
    t = (text or "").lower()
    if "egitim" in t or "education" in t:
        adim = {"eylem": "bolum_sil", "hedef": "EGITIM", "bolum": "",
                "alan": "", "deger": "", "guven": 0.95}
    elif "deneyim" in t:
        adim = {"eylem": "bolum_sil", "hedef": "IS DENEYIMI", "bolum": "",
                "alan": "", "deger": "", "guven": 0.95}
    else:
        adim = {"eylem": "belirsiz", "hedef": "", "bolum": "",
                "alan": "", "deger": "", "guven": 0.0}
    return {"adimlar": [adim]}, {"sure": 0.0, "thinking": ""}


# --- yardimcilar ------------------------------------------------------

def yukle(client, dosya, ad=""):
    with open(dosya, "rb") as f:
        return client.post("/upload",
                           files={"dosya": (Path(dosya).name, f.read(),
                                            "application/octet-stream")},
                           data={"ad": ad} if ad else {})


def bolumler(client, oid):
    r = client.get("/state", params={"id": oid}).json()
    return [s["heading"] for s in r.get("sections", [])]


# --- olcumler ---------------------------------------------------------

def test_uclar(app, client):
    print("\n[1] Uc noktalar")

    yollar = {getattr(r, "path", "") for r in app.app.routes}
    for yol in ("/upload", "/oturum", "/oturum/sec", "/oturum/sil",
                "/oturum/ad"):
        kontrol("uc var: {}".format(yol), yol in yollar, sorted(yollar))

    for yol in ("/", "/render", "/preview", "/state", "/command", "/undo",
                "/design/reset", "/health", "/voice"):
        kontrol("eski uc duruyor: {}".format(yol), yol in yollar)


def test_bos(client):
    print("\n[2] Hic oturum yokken cokmemeli")

    r = client.get("/oturum")
    kontrol("/oturum 200", r.status_code == 200, r.status_code)
    kontrol("/oturum bos liste", r.json().get("oturumlar") == [], r.json())
    kontrol("/oturum aktif bos", r.json().get("aktif", "x") == "", r.json())

    r = client.get("/state")
    kontrol("/state 200", r.status_code == 200, r.status_code)
    kontrol("/state ok False", r.json().get("ok") is False, r.json())

    r = client.post("/command", json={"text": "egitim bolumunu sil"})
    kontrol("/command 200", r.status_code == 200, r.status_code)
    kontrol("/command ok False", r.json().get("ok") is False, r.json())

    r = client.post("/undo")
    kontrol("/undo 200", r.status_code == 200, r.status_code)

    r = client.get("/render")
    kontrol("/render 200", r.status_code == 200, r.status_code)
    kontrol("/render ok False", r.json().get("ok") is False, r.json())

    r = client.get("/preview")
    kontrol("/preview cokmedi", r.status_code in (200, 404), r.status_code)

    r = client.get("/health")
    kontrol("/health 200", r.status_code == 200, r.status_code)
    kontrol("/health oturum sayisi 0", r.json().get("oturum") == 0, r.json())


def test_yukleme(client, tmp):
    print("\n[3] PDF yukleme")

    r = yukle(client, tmp / "ornek.pdf", ad="Birinci")
    kontrol("/upload 200", r.status_code == 200, r.status_code)
    d = r.json()
    kontrol("ok True", d.get("ok") is True, str(d)[:200])
    kontrol("ilk oturum 0001", d.get("id") == "0001", d.get("id"))
    kontrol("tur pdf", d.get("tur") == "pdf", d.get("tur"))
    kontrol("parca sayisi", (d.get("parca") or 0) >= 8, d.get("parca"))
    kontrol("bolum 2", d.get("bolum") == 2, d.get("bolum"))
    kontrol("kayit 2", d.get("kayit") == 2, d.get("kayit"))
    kontrol("kapsama %100", d.get("kapsama") == 1.0, d.get("kapsama"))
    kontrol("HTML uzunlugu", (d.get("uzunluk") or 0) > 1000, d.get("uzunluk"))

    dizin = Path(session.KOK) / "0001"
    for ad in ("cv_layout.json", "cv_structured.json", "cv_generated.html",
               "meta.json"):
        kontrol("0001/{} yazildi".format(ad), (dizin / ad).exists())
    kontrol("cv_page1.png yazildi", (dizin / "cv_page1.png").exists())
    kontrol("history klasoru var", (dizin / "history").is_dir())

    kontrol("uploads gecici klasore yazdi",
            any(Path(pipeline.UPLOADS).glob("ornek*.pdf")),
            list(Path(pipeline.UPLOADS).glob("*")))
    kontrol("ad form alani kullanildi",
            session.meta("0001").get("ad") == "Birinci",
            session.meta("0001").get("ad"))

    print("\n[4] DOCX yukleme, ikinci oturum")

    r = yukle(client, tmp / "ornek.docx")
    d = r.json()
    kontrol("ok True", d.get("ok") is True, str(d)[:200])
    kontrol("ikinci oturum 0002", d.get("id") == "0002", d.get("id"))
    kontrol("tur docx", d.get("tur") == "docx", d.get("tur"))
    kontrol("bolum 2", d.get("bolum") == 2, d.get("bolum"))
    kontrol("aktif oturum 0002", session.aktif() == "0002", session.aktif())
    kontrol("0001 dokunulmadi",
            (Path(session.KOK) / "0001" / "cv_structured.json").exists())


def test_reddetme(client, tmp):
    print("\n[5] Kabul edilmeyen dosya")

    kotu = tmp / "not.txt"
    kotu.write_text("merhaba", encoding="utf-8")
    r = yukle(client, kotu)
    kontrol("txt 200 dondu", r.status_code == 200, r.status_code)
    kontrol("txt ok False", r.json().get("ok") is False, r.json())
    kontrol("txt hata metni var", bool(r.json().get("error")), r.json())

    once = len(session.liste())
    r = yukle(client, kotu)
    kontrol("basarisiz yukleme oturum acmadi",
            len(session.liste()) == once, len(session.liste()))

    eski_mb = pipeline.MAX_MB
    try:
        pipeline.MAX_MB = 0.0001
        r = yukle(client, tmp / "ornek.pdf")
        kontrol("boyut siniri 200 dondu", r.status_code == 200, r.status_code)
        kontrol("boyut siniri ok False", r.json().get("ok") is False, r.json())
    finally:
        pipeline.MAX_MB = eski_mb

    bos = tmp / "bos.pdf"
    bos.write_bytes(b"")
    r = yukle(client, bos)
    kontrol("bos dosya cokmedi", r.status_code == 200, r.status_code)
    kontrol("bos dosya ok False", r.json().get("ok") is False, r.json())


def test_oturum_yonetimi(client):
    print("\n[6] Oturum yonetimi")

    r = client.get("/oturum").json()
    kontrol("liste 2 oturum", len(r.get("oturumlar", [])) == 2,
            len(r.get("oturumlar", [])))
    kontrol("aktif 0002", r.get("aktif") == "0002", r.get("aktif"))
    ilk = (r.get("oturumlar") or [{}])[0]
    for alan in ("id", "ad", "kaynak", "bolum", "hazir"):
        kontrol("ozet alani: {}".format(alan), alan in ilk, sorted(ilk))

    r = client.post("/oturum/sec", json={"id": "0001"})
    kontrol("sec 200", r.status_code == 200, r.status_code)
    kontrol("sec ok True", r.json().get("ok") is True, r.json())
    kontrol("aktif 0001 oldu", session.aktif() == "0001", session.aktif())

    r = client.post("/oturum/ad", json={"id": "0001", "ad": "Seyda CV"})
    kontrol("ad ok True", r.json().get("ok") is True, r.json())
    kontrol("ad diske yazildi", session.meta("0001").get("ad") == "Seyda CV",
            session.meta("0001").get("ad"))

    r = client.post("/oturum/sec", json={"id": "0009"})
    kontrol("olmayan oturum 200", r.status_code == 200, r.status_code)
    kontrol("olmayan oturum ok False", r.json().get("ok") is False, r.json())


def test_yalitim(client, tmp):
    print("\n[7] Oturum yalitimi: /command dogru klasore yaziyor mu")

    client.post("/oturum/sec", json={"id": "0001"})

    r = client.post("/command", json={"text": "egitim bolumunu sil",
                                      "id": "0001"})
    kontrol("/command 200", r.status_code == 200, r.status_code)
    d = r.json()
    kontrol("/command ok True", d.get("ok") is True, str(d)[:200])
    kontrol("/command applied", d.get("applied") is True, str(d)[:200])
    kontrol("/command id dondu", d.get("id") == "0001", d.get("id"))

    kontrol("0001 bolum 1 kaldi", bolumler(client, "0001") == ["IS DENEYIMI"],
            bolumler(client, "0001"))
    kontrol("0002 etkilenmedi", len(bolumler(client, "0002")) == 2,
            bolumler(client, "0002"))

    kontrol("gecmis 0001 icinde",
            len(list((Path(session.KOK) / "0001" / "history").glob("*.json"))) == 1)
    kontrol("gecmis 0002 bos",
            len(list((Path(session.KOK) / "0002" / "history").glob("*.json"))) == 0)

    r = client.post("/command", json={"text": "is deneyimi bolumunu sil",
                                      "id": "0002"})
    kontrol("0002 komutu ok", r.json().get("ok") is True, str(r.json())[:200])
    kontrol("0002 bolum 1 kaldi", bolumler(client, "0002") == ["EGITIM"],
            bolumler(client, "0002"))
    kontrol("0001 hala 1 bolum", bolumler(client, "0001") == ["IS DENEYIMI"],
            bolumler(client, "0001"))

    print("\n[8] Geri al oturuma bagli")

    r = client.post("/undo", params={"id": "0001"})
    kontrol("/undo 200", r.status_code == 200, r.status_code)
    kontrol("/undo applied", r.json().get("applied") is True, r.json())
    kontrol("0001 iki bolume dondu", len(bolumler(client, "0001")) == 2,
            bolumler(client, "0001"))
    kontrol("0002 geri alinmadi", bolumler(client, "0002") == ["EGITIM"],
            bolumler(client, "0002"))

    r = client.post("/undo", params={"id": "0001"})
    kontrol("bos gecmiste applied False",
            r.json().get("applied") is False, r.json())


def test_render_preview(client):
    print("\n[9] /render ve /preview oturumlu")

    r = client.get("/render", params={"id": "0002"})
    kontrol("/render ok", r.json().get("ok") is True, str(r.json())[:200])
    kontrol("/render id dondu", r.json().get("id") == "0002", r.json().get("id"))

    html = Path(session.KOK) / "0002" / "cv_generated.html"
    kontrol("HTML 0002 icine yazildi", html.exists(), html)
    kontrol("kok output kirletilmedi",
            not (Path(session.OUT) / "cv_generated.html").exists())

    r = client.get("/preview", params={"id": "0001"})
    kontrol("/preview 200", r.status_code == 200, r.status_code)
    kontrol("/preview 0001 icerigi", "IS DENEYIMI" in r.text
            or "Aksakal" in r.text, r.text[:120])


def test_reload(client):
    print("\n[10] --reload sonrasi yol kaybolmamali")

    commands.STRUCT = Path("C:/olmayan/cv_structured.json")
    commands.OUT = Path("C:/olmayan")
    commands.HIST = Path("C:/olmayan/history")
    render_cv.HTML_OUT = Path("C:/olmayan/cv_generated.html")

    r = client.get("/state", params={"id": "0001"})
    kontrol("/state bozuk yoldan sonra ok",
            r.json().get("ok") is True, str(r.json())[:200])
    kontrol("commands.STRUCT yeniden baglandi",
            Path(commands.STRUCT) == Path(session.KOK) / "0001" / "cv_structured.json",
            commands.STRUCT)

    session._baglanan = ""
    r = client.get("/state", params={"id": "0002"})
    kontrol("baglanan sifirlanmis olsa da calisti",
            r.json().get("ok") is True, str(r.json())[:200])
    kontrol("/state id dondu", r.json().get("id") == "0002", r.json().get("id"))


def test_kotu_kimlik(client):
    print("\n[11] Kotu oturum kimligi")

    for kotu in ("../gizli", "abc", "1", "0001/../0002", "0001 "):
        r = client.get("/state", params={"id": kotu})
        kontrol("reddedildi: {!r}".format(kotu),
                r.status_code == 200 and r.json().get("ok") is False,
                "{} {}".format(r.status_code, str(r.json())[:80]))


def test_sil(client):
    print("\n[12] Oturum silme")

    r = client.post("/oturum/sil", json={"id": "0002"})
    kontrol("sil ok True", r.json().get("ok") is True, r.json())
    kontrol("klasor gitti", not (Path(session.KOK) / "0002").exists())
    kontrol("liste 1 oturum",
            len(client.get("/oturum").json().get("oturumlar") or []) == 1)
    kontrol("aktif 0001 oldu", session.aktif() == "0001", session.aktif())

    r = client.post("/oturum/sil", json={"id": "0009"})
    kontrol("olmayan silme ok False", r.json().get("ok") is False, r.json())


def test_model(client, tmp):
    print("\n[13] Gercek model turu (--model)")

    import classify
    import llm
    classify.classify = _gercek_classify
    pipeline.cozumle = _gercek_cozumle

    kaynak = UPLOADS / "cv.pdf"
    if not kaynak.exists():
        kontrol("uploads/cv.pdf var", False, str(kaynak))
        return

    hedef = tmp / "gercek.pdf"
    shutil.copy2(kaynak, hedef)

    try:
        r = yukle(client, hedef, ad="Gercek")
    except Exception as e:                                   # noqa: BLE001
        kontrol("model turu calisti", False,
                "{}: {}".format(type(e).__name__, str(e)[:150]))
        return

    d = r.json()
    kontrol("model turu ok", d.get("ok") is True, str(d)[:200])
    kontrol("bolum >= 3", (d.get("bolum") or 0) >= 3, d.get("bolum"))
    kontrol("kayit >= 5", (d.get("kayit") or 0) >= 5, d.get("kayit"))
    kontrol("kapsama >= %90", (d.get("kapsama") or 0) >= 0.90, d.get("kapsama"))
    kontrol("ad okundu", len(d.get("ad") or "") >= 3, d.get("ad"))
    print("  sure {} sn (model {})".format(d.get("sure"), d.get("model_sure")))


_gercek_cozumle = None
_gercek_classify = None


def main():
    global _gercek_cozumle, _gercek_classify
    modelli = "--model" in sys.argv

    print("=" * 60)
    print("[0] Modul")

    tmp = Path(tempfile.mkdtemp(prefix="cvupload_"))
    eski = (session.KOK, session.OUT, pipeline.UPLOADS,
            commands.OUT, commands.STRUCT, commands.HIST,
            render_cv.OUT, render_cv.STRUCT, render_cv.LAYOUT,
            render_cv.HTML_OUT)
    _gercek_cozumle = pipeline.cozumle

    try:
        session.KOK = tmp / "oturum"
        session.OUT = tmp / "output"
        pipeline.UPLOADS = tmp / "uploads"
        (tmp / "output").mkdir(parents=True, exist_ok=True)

        # Oturum baglanmadan bir uc cagrilirsa gercek output\ yerine
        # gecici klasore dokunsun.
        commands.OUT = tmp / "output"
        commands.STRUCT = tmp / "output" / "cv_structured.json"
        commands.HIST = tmp / "output" / "history"
        render_cv.OUT = tmp / "output"
        render_cv.STRUCT = tmp / "output" / "cv_structured.json"
        render_cv.LAYOUT = tmp / "output" / "cv_layout.json"
        render_cv.HTML_OUT = tmp / "output" / "cv_generated.html"

        ornek_pdf(tmp / "ornek.pdf")
        ornek_docx(tmp / "ornek.docx")

        try:
            import app
        except Exception as e:                               # noqa: BLE001
            print("  KALDI  app import edilemedi   {}: {}".format(
                type(e).__name__, str(e)[:200]))
            print("\n" + "-" * 60)
            print("Sonuc: 0 gecti, 1 kaldi   (temel olcum)")
            return 1
        print("  gecti  app import edildi")

        import classify
        _gercek_classify = classify.classify
        pipeline.cozumle = sahte_cozumle
        classify.classify = sahte_classify

        from fastapi.testclient import TestClient
        client = TestClient(app.app)

        grup(test_uclar, app, client)
        grup(test_bos, client)
        grup(test_yukleme, client, tmp)
        grup(test_reddetme, client, tmp)
        grup(test_oturum_yonetimi, client)
        grup(test_yalitim, client, tmp)
        grup(test_render_preview, client)
        grup(test_reload, client)
        grup(test_kotu_kimlik, client)
        grup(test_sil, client)

        if modelli:
            grup(test_model, client, tmp)
        else:
            print("\n[13] Model turu atlandi (--model ile calistir)")
    finally:
        if _gercek_cozumle:
            pipeline.cozumle = _gercek_cozumle
        if _gercek_classify:
            import classify
            classify.classify = _gercek_classify
        (session.KOK, session.OUT, pipeline.UPLOADS,
         commands.OUT, commands.STRUCT, commands.HIST,
         render_cv.OUT, render_cv.STRUCT, render_cv.LAYOUT,
         render_cv.HTML_OUT) = eski
        shutil.rmtree(tmp, ignore_errors=True)

    print("\n[14] Gercek klasorler")
    kontrol("session.KOK geri konuldu", session.KOK == OUT / "oturum",
            session.KOK)
    kontrol("pipeline.UPLOADS geri konuldu", pipeline.UPLOADS == UPLOADS,
            pipeline.UPLOADS)
    kontrol("commands.STRUCT geri konuldu",
            commands.STRUCT == OUT / "cv_structured.json", commands.STRUCT)
    kontrol("uploads kirletilmedi",
            not (UPLOADS / "ornek.pdf").exists() and
            not (UPLOADS / "ornek.docx").exists())

    print("\n" + "-" * 60)
    print("Sonuc: {} gecti, {} kaldi".format(gecen + 1, kalan))
    return 1 if kalan else 0


if __name__ == "__main__":
    sys.exit(main())
