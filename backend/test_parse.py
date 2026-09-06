"""parse_docx.py olcumu: DOCX -> layout sozlugu.

Model yok, ag yok, kullanici dosyasi gerekmez. Test kendi DOCX ornegini
gecici klasorde uretir ve ayristiriciyi ona karsi olcer. Sema, PDF
tarafinin (parse_cv.py) urettigi layout semasiyla ayni olmak zorunda;
analyze_cv.py ikisini de ayni sekilde okuyacak.

uploads klasorunde .docx varsa ek olarak onun sayilari da yazilir
(bilgi amacli, puana girmez).

Calistirma (backend klasorunden):
    .\\.venv\\Scripts\\python.exe test_parse.py
"""

import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
UPLOADS = ROOT / "uploads"

gecen = kalan = 0


def kontrol(ad, kosul, ayrinti=""):
    global gecen, kalan
    try:
        ok = bool(kosul)
    except Exception as e:                                   # noqa: BLE001
        ok, ayrinti = False, "{}: {}".format(type(e).__name__, e)
    if ok:
        gecen += 1
        print("  gecti  {}".format(ad))
    else:
        kalan += 1
        print("  KALDI  {}   {}".format(ad, ayrinti))


def ornek_docx(path):
    """Basliklari, kalin/normal karisik satirlari ve maddeleri olan CV."""
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor

    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)      # A4, sablon Letter gelir
    s.left_margin = Mm(20)

    p = d.add_paragraph()
    r = p.add_run("Ayse Seyda Aksakal")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    p = d.add_paragraph()
    p.add_run("Yazilim Gelistirici").font.size = Pt(12)

    d.add_paragraph("")                       # bos paragraf, atlanmali

    p = d.add_paragraph()
    r = p.add_run("IS DENEYIMI")
    r.bold = True
    r.font.size = Pt(14)

    p = d.add_paragraph()
    r = p.add_run("Kidemli Muhendis")         # ayni paragrafta iki run
    r.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run("   2021 - 2024")
    r2.font.size = Pt(10)

    d.add_paragraph("Olcum altyapisini kurdu ve gecikmeyi dusurdu.",
                    style="List Bullet")
    d.add_paragraph("Ekip icin dagitim hattini yeniden yazdi.",
                    style="List Bullet")

    p = d.add_paragraph()
    r = p.add_run("EGITIM")
    r.bold = True
    r.font.size = Pt(14)

    d.add_paragraph("Anadolu Universitesi, Bilgisayar Muhendisligi")

    t = d.add_table(rows=1, cols=2)           # tablodaki metin de gelmeli
    t.rows[0].cells[0].text = "E-posta"
    t.rows[0].cells[1].text = "ornek@ornek.com"

    d.save(str(path))


def spans(data):
    return [s for p in data.get("pages", []) for s in p.get("spans", [])]


def metin(data):
    return " | ".join(s.get("text", "") for s in spans(data))


def test_sema(data, json_path):
    print("\n[1] Sema")

    kontrol("kok anahtarlar",
            all(k in data for k in ("file", "pages", "fonts_used", "colors_used")),
            sorted(data.keys()))

    pg = (data.get("pages") or [{}])[0]
    kontrol("sayfa anahtarlari",
            all(k in pg for k in ("page", "width_pt", "height_pt", "spans", "shapes")),
            sorted(pg.keys()))

    kontrol("shapes listesi var", isinstance(pg.get("shapes"), list))

    sp = spans(data)
    kontrol("parca uretildi", len(sp) >= 10, "{} parca".format(len(sp)))

    alanlar = ("text", "font", "size", "color", "bold", "bbox")
    eksik = [k for k in alanlar if sp and k not in sp[0]]
    kontrol("parca alanlari tam", sp and not eksik, "eksik: {}".format(eksik))

    kontrol("bbox 4 sayi",
            sp and all(len(s["bbox"]) == 4
                       and all(isinstance(v, (int, float)) for v in s["bbox"])
                       for s in sp))

    kontrol("bbox sifir alan degil",
            sp and all(s["bbox"][2] > s["bbox"][0] and s["bbox"][3] > s["bbox"][1]
                       for s in sp))

    kontrol("size sayi", sp and all(isinstance(s["size"], (int, float)) and s["size"] > 0
                                    for s in sp))

    kontrol("bold bool", sp and all(isinstance(s["bold"], bool) for s in sp))

    kontrol("renk #rrggbb",
            sp and all(isinstance(s["color"], str) and len(s["color"]) == 7
                       and s["color"][0] == "#" for s in sp),
            sp[0]["color"] if sp else "")

    kontrol("JSON diske yazildi", json_path.exists(), str(json_path))

    if json_path.exists():
        okundu = json.loads(json_path.read_text(encoding="utf-8"))
        kontrol("JSON tekrar okunabiliyor",
                len(spans(okundu)) == len(sp),
                "{} / {}".format(len(spans(okundu)), len(sp)))
    else:
        kontrol("JSON tekrar okunabiliyor", False)


def test_icerik(data):
    print("\n[2] Icerik")

    t = metin(data)
    sp = spans(data)

    kontrol("ad geldi", "Ayse Seyda Aksakal" in t)
    kontrol("bolum basliklari geldi", "IS DENEYIMI" in t and "EGITIM" in t)
    kontrol("madde metni geldi", "gecikmeyi dusurdu" in t)
    kontrol("tablo metni geldi", "ornek@ornek.com" in t)
    kontrol("bos paragraf atlandi",
            all(s.get("text", "").strip() for s in sp))
    kontrol("bosluk kirpildi",
            all(s["text"] == s["text"].strip() for s in sp))
    kontrol("ikinci run ayri parca", "2021 - 2024" in t)


def test_bicim(data):
    print("\n[3] Bicim")

    sp = spans(data)
    ad = next((s for s in sp if "Ayse Seyda" in s["text"]), None)
    govde = next((s for s in sp if "gecikmeyi dusurdu" in s["text"]), None)
    baslik = next((s for s in sp if s["text"].strip() == "IS DENEYIMI"), None)

    kontrol("ad kalin", ad and ad["bold"] is True)
    kontrol("ad puntosu buyuk", ad and ad["size"] >= 18, ad["size"] if ad else "-")
    kontrol("bolum basligi kalin", baslik and baslik["bold"] is True)
    kontrol("govde puntosu kucuk",
            govde and baslik and govde["size"] < baslik["size"],
            "{} < {}".format(govde["size"] if govde else "-",
                             baslik["size"] if baslik else "-"))
    kontrol("ad rengi korundu", ad and ad["color"].lower() == "#1a3c6e",
            ad["color"] if ad else "-")
    kontrol("font adi bos degil", sp and all(s["font"] for s in sp))
    kontrol("fonts_used dolu", len(data.get("fonts_used") or {}) >= 1)
    kontrol("colors_used dolu", len(data.get("colors_used") or {}) >= 1)


def test_duzen(data):
    print("\n[4] Duzen")

    pg = (data.get("pages") or [{}])[0]
    sp = spans(data)

    kontrol("A4 genislik", abs((pg.get("width_pt") or 0) - 595) <= 6,
            pg.get("width_pt"))
    kontrol("A4 yukseklik", abs((pg.get("height_pt") or 0) - 842) <= 6,
            pg.get("height_pt"))

    ys = [s["bbox"][1] for s in sp]
    kontrol("y yukaridan asagi artiyor", ys == sorted(ys))

    ad = next((s for s in sp if "Ayse Seyda" in s["text"]), None)
    son = sp[-1] if sp else None
    kontrol("ilk satir en ustte", ad and son and ad["bbox"][1] < son["bbox"][1])

    unvan = next((s for s in sp if "Kidemli Muhendis" in s["text"]), None)
    tarih = next((s for s in sp if "2021 - 2024" in s["text"]), None)
    kontrol("ayni paragraf ayni y",
            unvan and tarih and abs(unvan["bbox"][1] - tarih["bbox"][1]) < 1.0)
    kontrol("ikinci run daha sagda",
            unvan and tarih and tarih["bbox"][0] > unvan["bbox"][0])

    madde = next((s for s in sp if "gecikmeyi dusurdu" in s["text"]), None)
    kontrol("madde girintili",
            madde and ad and madde["bbox"][0] > ad["bbox"][0],
            "{} > {}".format(madde["bbox"][0] if madde else "-",
                             ad["bbox"][0] if ad else "-"))

    kontrol("sayfa icinde kaldi",
            all(0 <= s["bbox"][0] and s["bbox"][2] <= (pg.get("width_pt") or 0) + 1
                for s in sp))


def test_analyze_uyumu(data):
    print("\n[5] analyze_cv.py uyumu")

    sp = spans(data)
    try:
        satir = "y={} x={} {}pt {} | {}".format(
            sp[0]["bbox"][1], sp[0]["bbox"][0], sp[0]["size"],
            "KALIN" if sp[0]["bold"] else "normal", sp[0]["text"])
        ok = True
    except Exception as e:                                   # noqa: BLE001
        satir, ok = str(e), False
    kontrol("model satiri kurulabiliyor", ok, satir[:60])

    kontrol("tek sayfa dondu", len(data.get("pages") or []) == 1,
            len(data.get("pages") or []))
    kontrol("file alani dosya adi",
            str(data.get("file", "")).lower().endswith(".docx"), data.get("file"))


def gercek_dosya(parse):
    dosyalar = sorted(UPLOADS.glob("*.docx")) if UPLOADS.exists() else []
    if not dosyalar:
        print("\n[6] uploads klasorunde .docx yok, atlandi.")
        return
    tmp = Path(tempfile.mkdtemp(prefix="cvparse_"))
    try:
        for f in dosyalar[:3]:
            try:
                d = parse(f, tmp)
                sp = spans(d)
                print("\n[6] {} -> {} parca, {} font".format(
                    f.name, len(sp), len(d.get("fonts_used") or {})))
                for s in sp[:5]:
                    print("     {:>5}pt {} y={} | {}".format(
                        s["size"], s["color"], s["bbox"][1], s["text"][:40]))
            except Exception as e:                           # noqa: BLE001
                print("\n[6] {} -> HATA {}: {}".format(f.name, type(e).__name__, e))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    print("=" * 60)
    print("[0] Modul")
    try:
        import parse_docx
    except ImportError as e:
        print("  KALDI  parse_docx bulunamadi   {}".format(e))
        print("\n" + "-" * 60)
        print("Sonuc: 0 gecti, 1 kaldi   (temel olcum)")
        return 1

    parse = getattr(parse_docx, "parse", None)
    if not callable(parse):
        print("  KALDI  parse_docx.parse(path, out_dir) yok")
        print("\n" + "-" * 60)
        print("Sonuc: 0 gecti, 1 kaldi   (temel olcum)")
        return 1
    print("  gecti  parse_docx.parse hazir")

    tmp = Path(tempfile.mkdtemp(prefix="cvparse_"))
    try:
        docx_path = tmp / "ornek_cv.docx"
        ornek_docx(docx_path)
        print("  gecti  ornek DOCX uretildi ({} bayt)".format(
            docx_path.stat().st_size))

        data = parse(docx_path, tmp)
        if not isinstance(data, dict):
            print("  KALDI  parse() sozluk dondurmedi:", type(data).__name__)
            print("\n" + "-" * 60)
            print("Sonuc: 2 gecti, 1 kaldi")
            return 1

        json_path = tmp / "ornek_cv_layout.json"
        test_sema(data, json_path)
        test_icerik(data)
        test_bicim(data)
        test_duzen(data)
        test_analyze_uyumu(data)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    gercek_dosya(parse)

    print("\n" + "-" * 60)
    print("Sonuc: {} gecti, {} kaldi".format(gecen + 2, kalan))
    return 1 if kalan else 0


if __name__ == "__main__":
    sys.exit(main())
